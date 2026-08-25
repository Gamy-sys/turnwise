"""Small regression test for the Japanese four-line transcript contract."""
from __future__ import annotations

import tempfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document

from app.config import Settings
from app.exports import export_japanese_docx, export_txt
from app.models import DocumentMeta, Speaker, Token, Transcript, Turn
from app.pipeline.asr import ASRWord, _segment_japanese_words
from app.pipeline.japanese import populate_japanese_layers


def main() -> None:
    pieces = [
        ASRWord(text, i * 0.1, (i + 1) * 0.1)
        for i, text in enumerate(
            ["そ", "こ", "が", "もう", "少", "し、", "例えば", "お", "友", "達"]
        )
    ]
    segmented = _segment_japanese_words(pieces)
    segmented_text = [w.text for w in segmented]
    assert "そこ" in segmented_text
    assert "少し、" in segmented_text
    assert "お友達" in segmented_text

    turn = Turn(
        id="turn0",
        speaker="CM",
        start=0.0,
        end=2.0,
        tokens=[
            Token(
                id="w0",
                text="[で   そこ こ:   が:: もう 少し (.) たとえば (.)",
                start=0.0,
                end=2.0,
                speaker="CM",
            )
        ],
    )
    tr = Transcript(
        meta=DocumentMeta(project_id="ja-test", filename="sample.wav", duration=2.0),
        speakers=[Speaker(id="CM", label="CM")],
        turns=[turn],
        layout="japanese_four_line",
    )
    settings = Settings()
    settings.transcript_layout = "japanese_four_line"
    settings.japanese_auto_translate = False
    populate_japanese_layers(tr, settings)

    assert tr.turns[0].japanese
    assert tr.turns[0].japanese.transliteration.startswith("[de")
    assert "(.)" in tr.turns[0].japanese.transliteration

    # Structured translation parsing + marker preservation without a real API.
    settings.enable_openai = True
    settings.openai_api_key = "test-only"
    settings.japanese_auto_translate = True
    fake_json = """{"turns":[{"id":"turn0",
      "transliteration":"[de soko ko: ga:: mou sukoshi (.) tatoeba (.)",
      "direct_translation":"[then there SP more bit (.) for example (.)",
      "natural_translation":"[And if that point could be a bit more, for example, (.)"}]}"""
    fake_response = SimpleNamespace(
        choices=[SimpleNamespace(message=SimpleNamespace(content=fake_json))]
    )
    fake_client = SimpleNamespace(
        chat=SimpleNamespace(
            completions=SimpleNamespace(create=lambda **_: fake_response)
        )
    )
    with patch("openai.OpenAI", return_value=fake_client):
        populate_japanese_layers(tr, settings)
    assert tr.turns[0].japanese.direct_translation.startswith("[then")
    assert tr.turns[0].japanese.natural_translation.startswith("[And")

    tr.turns[0].japanese.direct_translation = (
        "[then there SP already more bit (.) for example (.)"
    )
    tr.turns[0].japanese.natural_translation = (
        "[And if that point could be a bit more, for example,"
    )

    with tempfile.TemporaryDirectory() as td:
        root = Path(td)
        txt = export_txt(tr, root / "japanese.txt")
        lines = txt.read_text(encoding="utf-8").splitlines()
        assert len(lines) == 4
        assert lines[0].startswith("1   CM  :    ")
        assert lines[1].strip().startswith("[de")
        assert lines[2].strip().startswith("[then")
        assert lines[3].strip().startswith("[And")

        docx = export_japanese_docx(tr, root / "japanese.docx")
        doc = Document(docx)
        assert len(doc.paragraphs) == 4
        assert doc.paragraphs[1].runs[0].italic is True
        assert doc.paragraphs[2].runs[0].bold is not True
        assert doc.paragraphs[3].runs[0].bold is True
        assert str(doc.paragraphs[3].runs[0].font.color.rgb) == "0070C0"

    print("JAPANESE FOUR-LINE TEST PASSED")


if __name__ == "__main__":
    main()


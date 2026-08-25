"""Export a Transcript to Jefferson .txt, Praat .TextGrid and ELAN .eaf.

Researchers live in Praat/ELAN, so these interop formats make the tool useful
beyond this app.
"""
from __future__ import annotations

import html
import math
import re
from pathlib import Path

from .models import Transcript


def render_japanese_four_line_text(tr: Transcript) -> str:
    """Render Japanese CA + transliteration + literal + natural English rows."""
    from .pipeline.ca import format_ca_line, render_turn_body

    width = len(str(max(1, len(tr.turns))))
    lines: list[str] = []
    for no, turn in enumerate(tr.turns, start=1):
        body = render_turn_body(turn)
        lines.append(format_ca_line(no, turn.speaker, body, num_width=width))
        if not turn.speaker:
            continue
        gutter = format_ca_line(no, turn.speaker, "", num_width=width)
        indent = " " * len(gutter)
        layers = turn.japanese
        lines.append(indent + (layers.transliteration if layers else ""))
        lines.append(indent + (layers.direct_translation if layers else ""))
        lines.append(indent + (layers.natural_translation if layers else ""))
    return "\n".join(lines)


def export_txt(tr: Transcript, dst: Path) -> Path:
    """Canonical Jefferson transcript (numbered CA layout)."""
    if tr.layout == "japanese_four_line":
        dst.write_text(render_japanese_four_line_text(tr).rstrip() + "\n", encoding="utf-8")
        return dst
    from .pipeline.ca import render_jefferson
    text = (tr.jefferson or "").strip()
    # Re-render when missing or still in the legacy "A:\\t..." form.
    if (not text) or (":\t" in text) or not re.match(r"^\s*\d+ {3}", text):
        text = render_jefferson(tr)
    dst.write_text(text + "\n", encoding="utf-8")
    return dst


def export_japanese_docx(tr: Transcript, dst: Path) -> Path:
    """Export the supplied-example four-line research layout to Word."""
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml.ns import qn
    from docx.shared import Inches, Mm, Pt, RGBColor

    from .pipeline.ca import format_ca_line, render_turn_body

    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)

    style_name = "Japanese CA Four Line"
    try:
        style = doc.styles[style_name]
    except KeyError:
        style = doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = "Courier New"
    style.font.size = Pt(10)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "ＭＳ 明朝")
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)

    def add_line(text: str, *, italic=False, bold=False, color="#000000", keep=True):
        p = doc.add_paragraph(style=style)
        p.paragraph_format.keep_with_next = keep
        run = p.add_run(text)
        run.font.name = "Courier New"
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "ＭＳ 明朝")
        run.font.size = Pt(10)
        run.italic = italic
        run.bold = bold
        try:
            run.font.color.rgb = RGBColor.from_string(color.lstrip("#").upper())
        except Exception:
            run.font.color.rgb = RGBColor(0, 0, 0)
        return p

    width = len(str(max(1, len(tr.turns))))
    for no, turn in enumerate(tr.turns, start=1):
        body = render_turn_body(turn)
        source = format_ca_line(no, turn.speaker, body, num_width=width)
        if not turn.speaker:
            add_line(source, keep=False)
            continue
        gutter = format_ca_line(no, turn.speaker, "", num_width=width)
        indent = " " * len(gutter)
        layers = turn.japanese
        add_line(source)
        add_line(indent + (layers.transliteration if layers else ""), italic=True)
        add_line(indent + (layers.direct_translation if layers else ""))
        add_line(
            indent + (layers.natural_translation if layers else ""),
            bold=True,
            color=layers.natural_color if layers else "#000000",
            keep=False,
        )

    doc.save(dst)
    return dst


_VBA_BAR = "=" * 34


def build_vba(tr: Transcript) -> str:
    """Plain, VBA-friendly timing export (word/sentence start+end in seconds).

    Word list carries no punctuation; punctuation lives only in the transcript.
    Words are listed strictly in chronological order across all speakers.
    """
    def is_word(t) -> bool:
        if not t.text.strip():
            return False
        return not any(c.type in ("laughter", "in_breath", "out_breath") for c in t.cues)

    # one sentence per spoken turn, ordered by time
    sentences = []
    for turn in tr.turns:
        if not turn.speaker:
            continue
        wt = [t for t in turn.tokens if is_word(t)]
        if not wt:
            continue
        last = wt[-1]
        ic = next((c for c in last.cues if c.type == "intonation"), None)
        term = "?" if (ic and "?" in (ic.symbol or "")) else "."
        text = " ".join(t.text.strip() for t in wt)
        sentences.append((wt[0].start, wt[-1].end, text, term, wt))
    sentences.sort(key=lambda s: s[0])

    words = []
    for s in sentences:
        for t in s[4]:
            words.append((t.start, t.end, t.text.strip()))
    words.sort(key=lambda w: (w[0], w[1]))

    def f(x: float) -> str:
        return f"{max(0.0, float(x)):.3f}"

    lines: list[str] = []
    lines += [_VBA_BAR, "TRANSCRIPT", _VBA_BAR, ""]
    lines += [" ".join(f"{s[2]}{s[3]}" for s in sentences)]
    lines += ["", _VBA_BAR, "WORD_TIMESTAMPS", _VBA_BAR, ""]
    lines += [f"{f(st)}|{f(en)}|{w}" for st, en, w in words]
    lines += ["", _VBA_BAR, "SENTENCE_TIMESTAMPS", _VBA_BAR, ""]
    lines += [f"{f(st)}|{f(en)}|{text}{term}" for st, en, text, term, _ in sentences]
    lines += ["", _VBA_BAR, "VBA_IMPORT", _VBA_BAR, ""]
    lines += [f"WordCount={len(words)}", ""]
    lines += [f"{f(st)}|{w}" for st, en, w in words]
    return "\n".join(lines) + "\n"


def export_vba(tr: Transcript, dst: Path) -> Path:
    dst.write_text(build_vba(tr), encoding="utf-8")
    return dst


# ---------------------------------------------------------------------------
# PowerPoint slides: the Jefferson transcript, chunked to fit on slides.
# ---------------------------------------------------------------------------

# Dark theme to match the app's transcript pane / karaoke video.
_SLIDE_BG = (0x0D, 0x11, 0x17)
_SLIDE_SPK = (0xF0, 0xC6, 0x74)   # amber speaker labels (accent2 in the UI)
_SLIDE_BODY = (0xE6, 0xED, 0xF3)  # light body text
_SLIDE_PAUSE = (0x8B, 0x94, 0x9A)  # muted for standalone pause lines


def _jefferson_lines(tr: Transcript) -> list[tuple[str, str]]:
    """Parse the rendered Jefferson transcript into (speaker, body) rows.

    Uses the exact rendered text so slides match the transcript pane (overlap
    brackets, elongation colons, pauses, latching, laughter, etc.).
    """
    from .pipeline.ca import parse_ca_line, render_jefferson

    text = tr.jefferson or ""
    if not text.strip():
        try:
            text = render_jefferson(tr)
        except Exception:
            text = ""
    # Upgrade legacy "A:\\t..." blobs so slides get the new spacing too.
    if text and (":\t" in text or not re.match(r"^\s*\d+ {3}", text)):
        try:
            text = render_jefferson(tr)
        except Exception:
            pass
    rows: list[tuple[str, str]] = []
    for raw in text.split("\n"):
        if not raw.strip():
            continue
        rows.append(parse_ca_line(raw))
    return rows


def build_transcript_pptx(tr: Transcript, dst: Path, font_size: int = 20) -> Path:
    """Lay the whole transcript out across slides at the given font size.

    Turns are packed greedily onto 16:9 slides using a monospace font so the CA
    alignment reads cleanly; long turns wrap and roll onto the next slide.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    margin, top, bottom = 0.5, 0.45, 0.45
    box_w = Inches(13.333 - 2 * margin)
    box_h = Inches(7.5 - top - bottom)

    # Estimate monospace metrics so we know how much fits per slide.
    char_w_in = 0.55 * font_size / 72.0
    line_h_in = 1.22 * font_size / 72.0
    usable_w_in = 13.333 - 2 * margin
    usable_h_in = 7.5 - top - bottom
    cpl = max(24, int(usable_w_in / char_w_in))          # chars per visual line
    lps = max(4, int(usable_h_in / (line_h_in * 1.12)))  # lines per slide (conservative)

    rows = _jefferson_lines(tr)
    if not rows:
        rows = [("", "(empty transcript)")]

    # Greedy pack rows into slides by estimated wrapped-line count. Each row
    # carries a sequential CA line number (nbsp-padded gutter).
    nwidth = len(str(max(1, len(rows))))

    def _numpfx(n: int) -> str:
        return f"{n}".rjust(nwidth, "\u00a0") + "\u00a0\u00a0\u00a0"

    pages: list[list[tuple[str, str, str]]] = []
    cur: list[tuple[str, str, str]] = []
    used = 0
    for n, (spk, body) in enumerate(rows, start=1):
        full_len = len(f"{spk}  :    {body}") if spk else len(f"  :    {body}")
        wl = max(1, math.ceil(full_len / cpl))
        if used + wl > lps and cur:
            pages.append(cur)
            cur, used = [], 0
        cur.append((_numpfx(n), spk, body))
        used += wl
    if cur:
        pages.append(cur)

    bg = RGBColor(*_SLIDE_BG)
    c_spk = RGBColor(*_SLIDE_SPK)
    c_body = RGBColor(*_SLIDE_BODY)
    c_pause = RGBColor(*_SLIDE_PAUSE)

    for page in pages:
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = bg

        tb = slide.shapes.add_textbox(Inches(margin), Inches(top), box_w, box_h)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP

        first = True
        for npfx, spk, body in page:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.space_after = Pt(max(4, font_size // 3))
            rn = p.add_run()
            rn.text = npfx
            rn.font.size = Pt(font_size)
            rn.font.name = "Consolas"
            rn.font.color.rgb = c_pause
            spfx = f"{spk or ''}\u00a0\u00a0:\u00a0\u00a0\u00a0\u00a0"
            if spk:
                r1 = p.add_run()
                r1.text = spfx
                r1.font.bold = True
                r1.font.size = Pt(font_size)
                r1.font.name = "Consolas"
                r1.font.color.rgb = c_spk
                r2 = p.add_run()
                r2.text = body
                r2.font.size = Pt(font_size)
                r2.font.name = "Consolas"
                r2.font.color.rgb = c_body
            else:
                rs = p.add_run()
                rs.text = spfx
                rs.font.size = Pt(font_size)
                rs.font.name = "Consolas"
                rs.font.color.rgb = c_pause
                r = p.add_run()
                r.text = body
                r.font.italic = True
                r.font.size = Pt(font_size)
                r.font.name = "Consolas"
                r.font.color.rgb = c_pause

    prs.save(str(dst))
    return dst


def export_textgrid(tr: Transcript, dst: Path) -> Path:
    """Minimal Praat TextGrid: one interval tier per speaker with word text."""
    xmax = max(tr.meta.duration, 0.001)
    speakers = tr.speakers or []
    tiers = []
    for spk in speakers:
        intervals = []
        for turn in tr.turns:
            if turn.speaker != spk.id:
                continue
            for tok in turn.tokens:
                intervals.append((tok.start, tok.end, tok.text))
        intervals.sort()
        tiers.append((spk.label, intervals))
    if not tiers:
        tiers = [("A", [])]

    lines = [
        'File type = "ooTextFile"',
        'Object class = "TextGrid"',
        "",
        "xmin = 0",
        f"xmax = {xmax}",
        "tiers? <exists>",
        f"size = {len(tiers)}",
        "item []:",
    ]
    for ti, (name, intervals) in enumerate(tiers, start=1):
        lines.append(f"    item [{ti}]:")
        lines.append('        class = "IntervalTier"')
        lines.append(f'        name = "{name}"')
        lines.append("        xmin = 0")
        lines.append(f"        xmax = {xmax}")
        # build contiguous intervals (fill gaps with empty labels)
        filled = []
        prev = 0.0
        for s, e, txt in intervals:
            if s > prev:
                filled.append((prev, s, ""))
            filled.append((s, e, txt))
            prev = e
        if prev < xmax:
            filled.append((prev, xmax, ""))
        if not filled:
            filled = [(0.0, xmax, "")]
        lines.append(f"        intervals: size = {len(filled)}")
        for ii, (s, e, txt) in enumerate(filled, start=1):
            lines.append(f"        intervals [{ii}]:")
            lines.append(f"            xmin = {s}")
            lines.append(f"            xmax = {e}")
            safe = txt.replace('"', '""')
            lines.append(f'            text = "{safe}"')
    dst.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return dst


def export_eaf(tr: Transcript, dst: Path) -> Path:
    """Minimal ELAN .eaf with one tier per speaker."""
    time_slots: dict[str, int] = {}
    order: list[tuple[str, int]] = []

    def slot(ms: int) -> str:
        key = str(ms)
        if key not in time_slots:
            sid = f"ts{len(time_slots) + 1}"
            time_slots[key] = sid
            order.append((sid, ms))
            return sid
        return time_slots[key]

    tier_annotations: dict[str, list[tuple[str, str, str]]] = {}
    ann_id = 1
    for turn in tr.turns:
        if not turn.speaker:
            continue
        tier_annotations.setdefault(turn.speaker, [])
        for tok in turn.tokens:
            a = slot(int(tok.start * 1000))
            b = slot(int(tok.end * 1000))
            tier_annotations[turn.speaker].append((a, b, tok.text))

    parts = ['<?xml version="1.0" encoding="UTF-8"?>']
    parts.append('<ANNOTATION_DOCUMENT AUTHOR="CA Studio" FORMAT="3.0" VERSION="3.0" '
                 'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">')
    parts.append(f'  <HEADER MEDIA_FILE="" TIME_UNITS="milliseconds">')
    parts.append(f'    <MEDIA_DESCRIPTOR MEDIA_URL="{html.escape(tr.meta.filename)}" '
                 f'MIME_TYPE="audio/x-wav"/>')
    parts.append("  </HEADER>")
    parts.append("  <TIME_ORDER>")
    for sid, ms in order:
        parts.append(f'    <TIME_SLOT TIME_SLOT_ID="{sid}" TIME_VALUE="{ms}"/>')
    parts.append("  </TIME_ORDER>")

    aid = 1
    for spk, anns in tier_annotations.items():
        parts.append(f'  <TIER LINGUISTIC_TYPE_REF="default-lt" TIER_ID="{html.escape(spk)}">')
        for a, b, txt in anns:
            parts.append(f'    <ANNOTATION>')
            parts.append(f'      <ALIGNABLE_ANNOTATION ANNOTATION_ID="a{aid}" '
                         f'TIME_SLOT_REF1="{a}" TIME_SLOT_REF2="{b}">')
            parts.append(f'        <ANNOTATION_VALUE>{html.escape(txt)}</ANNOTATION_VALUE>')
            parts.append(f'      </ALIGNABLE_ANNOTATION>')
            parts.append(f'    </ANNOTATION>')
            aid += 1
        parts.append("  </TIER>")
    parts.append('  <LINGUISTIC_TYPE GRAPHIC_REFERENCES="false" LINGUISTIC_TYPE_ID="default-lt" '
                 'TIME_ALIGNABLE="true"/>')
    parts.append("</ANNOTATION_DOCUMENT>")
    dst.write_text("\n".join(parts) + "\n", encoding="utf-8")
    return dst
